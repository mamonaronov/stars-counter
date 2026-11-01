"""Excel attendance journal, one sheet per period.

The sheet follows the class register: weekday and day of the month over each
lesson, the lesson title under them, and a weight row (1 for a mandatory
lesson). A student row has the course, name, percent, visit count, direction,
enrollment year, prayer track and Judaica group. A visit is 1. A free lesson
has no weight. An excused miss is also 1: pale green for a written leave, pale
blue for a medical excuse. A day of only free lessons is pale orange. Lessons
outside the dates a student is enrolled are pale cyan, and that student's sums
cover only the lessons that count for them. The percent column uses the same
pastel scale and is not listed in the color key.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from database.models import Event
from services.stats import STATUS_EXCUSED, STATUS_NOT_COUNTED, STATUS_PRESENT, EventRow, StudentStats

_THIN = Side(style="thin", color="000000")
_THICK = Side(style="thick", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FREE_FILL = PatternFill("solid", fgColor="FCE5CD")
_LEAVE_FILL = PatternFill("solid", fgColor="D9EAD3")
_SICK_FILL = PatternFill("solid", fgColor="C9DAF8")
_OUTSIDE_FILL = PatternFill("solid", fgColor="C9EDED")
_FONT = Font(name="Arial", size=10, color="000000")
_FONT_BOLD = Font(name="Arial", size=10, bold=True, color="000000")
_CENTER = Alignment(horizontal="center")
_LEFT = Alignment(horizontal="left")
_ROTATED = Alignment(horizontal="center", textRotation=90)
_ROTATED_LEFT = Alignment(horizontal="left", textRotation=90)
_ROTATED_PLAIN = Alignment(textRotation=90)

_WEEKDAYS = ("понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье")
_EXCUSE_FILL = {
    "leave": _LEAVE_FILL,
    "sick": _SICK_FILL,
}
def _cf_fill(hex_color: str) -> PatternFill:
    """Solid fill Excel will actually paint in a conditional format.

    A differential style needs the colour on both sides, with a set alpha.
    fgColor alone is stored as transparent and the cell stays blank.
    """
    color = f"FF{hex_color}"
    return PatternFill(patternType="solid", fgColor=color, bgColor=color)


# The cell holds a fraction (1 = 100%). Pale green from 85%, then orange, yellow, pink, red.
_PERCENT_BANDS = (
    ("0.85", _cf_fill("B6D7A8")),
    ("0.7", _cf_fill("F9CB9C")),
    ("0.6", _cf_fill("FFE599")),
    ("0.5", _cf_fill("F4CCCC")),
    ("0", _cf_fill("EA9999")),
)
_LEGEND = (
    (_LEAVE_FILL, "Освобождение по заявлению"),
    (_SICK_FILL, "Освобождение по мед причине"),
    (_FREE_FILL, "Свободный день"),
    (_OUTSIDE_FILL, "Вне срока ученика"),
)

_IDENTITY_HEADERS = (
    (1, "Курс", _ROTATED),
    (2, "ФИО", _CENTER),
    (3, "% посещения", _ROTATED_PLAIN),
    (4, "Кол-во посещения", _ROTATED),
    (5, "Направление", _ROTATED),
    (6, "Год поступления", _ROTATED_LEFT),
    (7, "Направление", _CENTER),
    (8, "Группа иудаики", _CENTER),
)
_FIRST_EVENT_COL = 9
_HEADER_ROW = 3
_WEIGHT_ROW = 4
_FIRST_STUDENT_ROW = 5
_BAD_SHEET = re.compile(r"[\\*?:/\[\]]")


@dataclass(frozen=True, slots=True)
class PeriodSection:
    period_title: str
    period_from: date | None
    period_to: date | None
    stats: list[StudentStats]
    events: list[Event]
    interim: bool = False


def build_report_xlsx(
    *,
    org_title: str,
    period_title: str,
    period_from: date | None,
    period_to: date | None,
    stats: list[StudentStats],
    events: list[Event],
    interim: bool = False,
) -> bytes:
    wb = Workbook()
    wb.properties.title = org_title
    ws = wb.active
    ws.title = _sheet_name(period_title, set())
    _write_journal(
        ws,
        period_from=period_from,
        period_to=period_to,
        stats=stats,
        events=events,
        interim=interim,
    )
    return _save(wb)


def build_periods_xlsx(*, org_title: str, sections: list[PeriodSection]) -> bytes:
    """One journal sheet for every period, in chronological order."""
    if not sections:
        raise ValueError("no sections")
    wb = Workbook()
    wb.properties.title = org_title
    used: set[str] = set()
    first = True
    for section in sections:
        if first:
            ws = wb.active
            ws.title = _sheet_name(section.period_title, used)
            first = False
        else:
            ws = wb.create_sheet(_sheet_name(section.period_title, used))
        _write_journal(
            ws,
            period_from=section.period_from,
            period_to=section.period_to,
            stats=section.stats,
            events=section.events,
            interim=section.interim,
        )
    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sheet_name(raw: str, used: set[str]) -> str:
    cleaned = _BAD_SHEET.sub(" ", raw)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" '") or "Лист"
    cleaned = cleaned[:31]
    candidate = cleaned
    number = 2
    while candidate.casefold() in used:
        suffix = f" {number}"
        candidate = f"{cleaned[: 31 - len(suffix)]}{suffix}".strip()
        number += 1
    used.add(candidate.casefold())
    return candidate


def _course_number(enrollment_year: int | None, on: date | None) -> int | None:
    """Course as of *on*. The academic year starts in September."""
    if enrollment_year is None or on is None:
        return None
    year = enrollment_year if enrollment_year >= 100 else 2000 + enrollment_year
    course = on.year - year + (1 if on.month >= 9 else 0)
    return course if course >= 1 else 1


def _short_year(enrollment_year: int | None) -> int | None:
    if enrollment_year is None:
        return None
    return enrollment_year % 100


def _day_groups(events: list[Event]) -> list[list[Event]]:
    groups: list[list[Event]] = []
    for event in events:
        if groups and groups[-1][0].day == event.day:
            groups[-1].append(event)
        else:
            groups.append([event])
    return groups


def _put(
    cell,
    value,
    *,
    bold: bool = False,
    align: Alignment | None = None,
    fmt: str | None = None,
    fill: PatternFill | None = None,
    border: bool = False,
) -> None:
    cell.value = value
    cell.font = _FONT_BOLD if bold else _FONT
    if align is not None:
        cell.alignment = align
    if fmt is not None:
        cell.number_format = fmt
    if fill is not None:
        cell.fill = fill
    if border:
        cell.border = _BORDER


def _write_journal(
    ws: Worksheet,
    *,
    period_from: date | None,
    period_to: date | None,
    stats: list[StudentStats],
    events: list[Event],
    interim: bool,
) -> None:
    if interim:
        _put(ws["A1"], "промежуточный", bold=True)

    _put(ws["H1"], "День недели", bold=True, align=_CENTER)
    _put(ws["H2"], "Число месяца", bold=True, align=_CENTER)
    for col, text, align in _IDENTITY_HEADERS:
        _put(ws.cell(row=_HEADER_ROW, column=col, value=text), text, bold=True, align=align, border=True)
    for col in range(1, 9):
        _put(ws.cell(row=_WEIGHT_ROW, column=col), None, bold=True, align=_CENTER, border=True)

    groups = _day_groups(events)
    day_free_by_id: dict[int, bool] = {}
    col = _FIRST_EVENT_COL
    for group in groups:
        start = col
        day_free = all(event.is_free for event in group)
        width = 12.63 if len(group) == 1 else 6.63
        for event in group:
            day_free_by_id[event.id] = day_free
            fill = _FREE_FILL if day_free else None
            _put(
                ws.cell(row=_HEADER_ROW, column=col),
                event.title,
                bold=True,
                align=_ROTATED,
                fill=fill,
                border=True,
            )
            weight = None if event.is_free else 1
            _put(
                ws.cell(row=_WEIGHT_ROW, column=col),
                weight,
                bold=True,
                align=_CENTER,
                fill=fill,
                border=True,
            )
            ws.column_dimensions[get_column_letter(col)].width = width
            col += 1
        end = col - 1
        header_fill = _FREE_FILL if day_free else None
        _put(
            ws.cell(row=1, column=start),
            _WEEKDAYS[group[0].day.weekday()],
            bold=True,
            align=_CENTER,
            fill=header_fill,
            border=True,
        )
        _put(
            ws.cell(row=2, column=start),
            group[0].day.day,
            bold=True,
            align=_CENTER,
            fill=header_fill,
            border=True,
        )
        for column in range(start, end + 1):
            for row in (1, 2):
                cell = ws.cell(row=row, column=column)
                cell.border = _BORDER
                if header_fill is not None:
                    cell.fill = header_fill
        if end > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)

    last_event_col = _FIRST_EVENT_COL + len(events) - 1 if events else 8
    weight_col = last_event_col + 1 if events else None
    weight_ref = f"${get_column_letter(weight_col)}${_WEIGHT_ROW}" if weight_col else None
    if weight_col is not None:
        first_letter = get_column_letter(_FIRST_EVENT_COL)
        last_letter = get_column_letter(last_event_col)
        _put(ws.cell(row=_HEADER_ROW, column=weight_col), None, border=True)
        _put(
            ws.cell(row=_WEIGHT_ROW, column=weight_col),
            f"=SUM({first_letter}{_WEIGHT_ROW}:{last_letter}{_WEIGHT_ROW})",
            bold=True,
            align=_CENTER,
            border=True,
        )
        ws.column_dimensions[get_column_letter(weight_col)].width = 7.88

    ordered = sorted(stats, key=lambda item: item.student.full_name.casefold())
    as_of = period_to or period_from

    for offset, item in enumerate(ordered):
        row = _FIRST_STUDENT_ROW + offset
        bounds = _counted_bounds(item.student, events)
        _put(ws.cell(row=row, column=1), _course_number(item.student.enrollment_year, as_of), align=_CENTER, border=True)
        _put(ws.cell(row=row, column=2), item.student.full_name, border=True)
        _put(ws.cell(row=row, column=4), _count_formula(row, events, bounds), border=True)
        _put(
            ws.cell(row=row, column=3),
            _percent_formula(row, item.mandatory_total, events, bounds, weight_ref),
            fmt="0%",
            border=True,
        )
        _put(ws.cell(row=row, column=5), item.student.direction, border=True)
        _put(
            ws.cell(row=row, column=6),
            _short_year(item.student.enrollment_year),
            align=_LEFT,
            border=True,
        )
        _put(ws.cell(row=row, column=7), item.student.track, border=True)
        _put(ws.cell(row=row, column=8), item.student.judaica_group, border=True)
        if weight_col is not None:
            _put(ws.cell(row=row, column=weight_col), None, border=True)
        by_event = {record.event.id: record for record in item.rows}
        for index, event in enumerate(events):
            column = _FIRST_EVENT_COL + index
            record = by_event.get(event.id)
            _put(
                ws.cell(row=row, column=column),
                _mark(record),
                align=_CENTER,
                fill=_mark_fill(record, day_free=day_free_by_id.get(event.id, False)),
                border=True,
            )

    last_student_row = _FIRST_STUDENT_ROW + len(ordered) - 1 if ordered else _WEIGHT_ROW
    filter_end = last_student_row
    if ordered and events:
        _write_totals(ws, ordered, events, day_free_by_id, last_student_row)
        filter_end = last_student_row + 2

    if ordered and any(item.percent is not None for item in ordered):
        percent_range = f"C{_FIRST_STUDENT_ROW}:C{last_student_row}"
        for formula, fill in _PERCENT_BANDS:
            ws.conditional_formatting.add(
                percent_range,
                CellIsRule(
                    operator="greaterThanOrEqual",
                    formula=[formula],
                    fill=fill,
                    stopIfTrue=True,
                ),
            )

    ws.row_dimensions[_HEADER_ROW].height = 120
    ws.column_dimensions["A"].width = 4.88
    ws.column_dimensions["B"].width = 29.13
    ws.column_dimensions["C"].width = 7.63
    ws.column_dimensions["D"].width = 4.88
    ws.column_dimensions["E"].width = 5.5
    ws.column_dimensions["F"].width = 4.63
    ws.column_dimensions["G"].width = 11.63
    ws.column_dimensions["H"].width = 21.13
    ws.auto_filter.ref = f"A{_WEIGHT_ROW}:{get_column_letter(last_event_col)}{filter_end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 0
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75


def _counted_bounds(student, events: list[Event]) -> tuple[int, int] | None:
    indexes = [index for index, event in enumerate(events) if student.counts_on(event.day)]
    if not indexes:
        return None
    return indexes[0], indexes[-1]


def _count_formula(row: int, events: list[Event], bounds: tuple[int, int] | None) -> str | int:
    if not events or bounds is None:
        return 0
    start, end = bounds
    first = get_column_letter(_FIRST_EVENT_COL + start)
    last = get_column_letter(_FIRST_EVENT_COL + end)
    return f"=SUM({first}{row}:{last}{row})"


def _percent_formula(
    row: int,
    mandatory: int,
    events: list[Event],
    bounds: tuple[int, int] | None,
    weight_ref: str | None,
) -> str:
    if mandatory <= 0 or not events or bounds is None:
        return "—"
    start, end = bounds
    if weight_ref is not None and start == 0 and end == len(events) - 1:
        return f"=D{row}/{weight_ref}"
    first = get_column_letter(_FIRST_EVENT_COL + start)
    last = get_column_letter(_FIRST_EVENT_COL + end)
    return f"=D{row}/SUM(${first}${_WEIGHT_ROW}:${last}${_WEIGHT_ROW})"


def _mark(record: EventRow | None) -> int | None:
    if record is None:
        return None
    if record.status in {STATUS_PRESENT, STATUS_EXCUSED}:
        return 1
    return None


def _mark_fill(record: EventRow | None, *, day_free: bool) -> PatternFill | None:
    if day_free:
        return _FREE_FILL
    if record is not None and record.status == STATUS_NOT_COUNTED:
        return _OUTSIDE_FILL
    if record is not None and record.status == STATUS_EXCUSED and record.excuse is not None:
        return _EXCUSE_FILL.get(record.excuse.kind)
    return None


def _write_totals(
    ws: Worksheet,
    stats: list[StudentStats],
    events: list[Event],
    day_free_by_id: dict[int, bool],
    last_student_row: int,
) -> None:
    percent_row = last_student_row + 1
    sum_row = last_student_row + 2
    count = len(stats)
    _write_legend(ws, percent_row)
    for index, event in enumerate(events):
        column = _FIRST_EVENT_COL + index
        letter = get_column_letter(column)
        fill = _FREE_FILL if day_free_by_id.get(event.id) else None
        _put(
            ws.cell(row=sum_row, column=column),
            f"=SUM({letter}{_FIRST_STUDENT_ROW}:{letter}{last_student_row})",
            align=_CENTER,
            fill=fill,
            border=True,
        )
        _put(
            ws.cell(row=percent_row, column=column),
            f"={letter}{sum_row}/{count}",
            align=_CENTER,
            fmt="0.00%",
            fill=fill,
            border=True,
        )


def _write_legend(ws: Worksheet, start_row: int) -> None:
    """Color key in F:H, from the summary rows downward, as in the class register."""
    last = len(_LEGEND) - 1
    for index, (fill, text) in enumerate(_LEGEND):
        row = start_row + index
        top = _THICK if index == 0 else None
        bottom = _THICK if index == last else None
        swatch = ws.cell(row=row, column=6)
        _put(swatch, None, bold=index == 0, fill=fill)
        swatch.border = Border(left=_THICK, top=top, bottom=bottom)
        label = ws.cell(row=row, column=7)
        _put(label, text)
        label.border = Border(top=top, bottom=bottom)
        edge = ws.cell(row=row, column=8)
        _put(edge, None)
        edge.border = Border(top=top, right=_THICK, bottom=bottom)
