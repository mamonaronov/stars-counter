from __future__ import annotations

from dataclasses import replace
from io import BytesIO

from openpyxl import load_workbook

from config import Config
from database.queries import Repo
from services.group_export import KEY_GROUP, markers_chat_id, prepare_markers_export, publish_markers_export


class _Msg:
    def __init__(self, message_id: int) -> None:
        self.message_id = message_id


class FakeBot:
    def __init__(self) -> None:
        self.docs: list[dict] = []
        self.deleted: list[int] = []
        self.pinned: list[int] = []
        self._n = 0

    async def send_document(self, chat_id, document, caption=None, disable_notification=None):
        self._n += 1
        self.docs.append(
            {
                "chat_id": chat_id,
                "caption": caption,
                "silent": disable_notification,
                "filename": document.filename,
                "data": document.data,
            }
        )
        return _Msg(self._n)

    async def pin_chat_message(self, chat_id, message_id, disable_notification=None):
        self.pinned.append(message_id)

    async def delete_message(self, chat_id, message_id):
        self.deleted.append(message_id)


async def test_export_follows_attendance_and_replaces_the_previous_file(repo: Repo, config: Config):
    cfg = replace(config, markers_group_id=-100123)
    bot = FakeBot()
    assert await publish_markers_export(bot, repo, cfg, silent=True) == "empty"
    assert bot.docs == []

    await repo.create_student("Антонов Антон", 2024, "2026-09-01", 1)
    first = await repo.create_event("Лекция", "mandatory", "2026-09-01", "10:00", None, 1)
    await repo.update_event(first.id, closes_month=1, month_title="Сентябрь 2026")
    current = await repo.create_event("Семинар", "mandatory", "2026-10-02", "12:00", None, 1)

    prepared = await prepare_markers_export(repo, cfg)
    assert prepared.payload is not None
    book = load_workbook(BytesIO(prepared.payload.data))
    assert len(book.sheetnames) == 4
    assert "промежуточный" in book[book.sheetnames[2]]["A1"].value

    assert await publish_markers_export(bot, repo, cfg, silent=True) == "sent"
    assert await publish_markers_export(bot, repo, cfg, silent=True) == "unchanged"
    assert len(bot.docs) == 1
    assert bot.docs[0]["chat_id"] == -100123
    assert bot.docs[0]["silent"] is True
    assert bot.pinned == [1]
    assert "Посещаемость" in bot.docs[0]["caption"]

    student = (await repo.list_students())[0]
    await repo.set_presence(current.id, student.id, True, 1)
    assert await publish_markers_export(bot, repo, cfg, silent=False) == "sent"
    assert len(bot.docs) == 2
    assert bot.deleted == [1]
    assert bot.docs[1]["silent"] is False

    await repo.db.set_system(KEY_GROUP, "-200555")
    assert await markers_chat_id(repo.db, cfg) == -200555
    assert await publish_markers_export(bot, repo, cfg, silent=True, force=True) == "sent"
    assert bot.docs[-1]["chat_id"] == -200555


async def test_export_stays_off_without_a_group(repo: Repo, config: Config):
    bot = FakeBot()
    await repo.create_student("Антонов Антон", 2024, "2026-09-01", 1)
    await repo.create_event("Лекция", "mandatory", "2026-09-01", None, None, 1)
    assert await publish_markers_export(bot, repo, config, silent=True) == "unconfigured"
    assert bot.docs == []
