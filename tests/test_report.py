from __future__ import annotations

import re
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from services.report import PeriodSection, build_periods_xlsx, build_report_xlsx
from services.stats import compute_group_stats
from tests.conftest import event, excuse, student


def _book(data: bytes):
    return load_workbook(BytesIO(data))


def test_report_journal_layout():
    students = [
        student(1, "Антонов Антон", direction="ПИН", judaica_group="рав Исроэль Фельдшеров"),
        student(2, "Борисов Борис"),
    ]
    events = [event(1, "2026-09-01"), event(2, "2026-09-03", kind="free", title="Доп"), event(3, "2026-09-05")]
    stats = compute_group_stats(students, events, {1: {1, 2}, 2: {1}}, [excuse(2, "2026-09-05", "2026-09-05")])
    wb = _book(
        build_report_xlsx(
            org_title="Посещаемость групп",
            period_title="Сентябрь 2026",
            period_from=date(2026, 9, 1),
            period_to=date(2026, 9, 5),
            stats=stats,
            events=events,
        )
    )
    ws = wb["Сентябрь 2026"]
    assert ws["H1"].value == "День недели"
    assert ws["H2"].value == "Число месяца"
    assert ws["A3"].value == "Курс"
    assert ws["B3"].value == "ФИО"
    assert ws["C3"].value == "% посещения"
    assert ws["D3"].value == "Кол-во посещения"
    assert ws["F3"].value == "Год поступления"
    assert ws["H3"].value == "Группа иудаики"
    assert (ws["I1"].value, ws["J1"].value, ws["K1"].value) == ("вторник", "четверг", "суббота")
    assert (ws["I2"].value, ws["J2"].value, ws["K2"].value) == (1, 3, 5)
    assert (ws["I3"].value, ws["J3"].value, ws["K3"].value) == ("Лекция", "Доп", "Лекция")
    assert (ws["I4"].value, ws["J4"].value, ws["K4"].value) == (1, None, 1)
    assert ws["L4"].value == "=SUM(I4:K4)"
    assert ws["J3"].fill.fgColor.rgb.endswith("D9EAD3")

    assert ws["B5"].value == "Антонов Антон"
    assert ws["A5"].value == 3
    assert ws["E5"].value == "ПИН"
    assert ws["F5"].value == 24
    assert ws["H5"].value == "рав Исроэль Фельдшеров"
    assert ws["E6"].value is None
    assert ws["H6"].value is None
    assert ws["C5"].value == "=D5/$L$4"
    assert ws["D5"].value == "=SUM(I5:K5)"
    assert (ws["I5"].value, ws["J5"].value, ws["K5"].value) == (1, 1, None)

    assert ws["B6"].value == "Борисов Борис"
    assert ws["C6"].value == "=D6/$L$4"
    assert (ws["I6"].value, ws["J6"].value, ws["K6"].value) == (1, None, 1)
    assert "мед" in ws["K6"].comment.text
    assert ws["B7"].value == "Доля по занятию"
    assert ws["I8"].value == "=SUM(I5:I6)"
    assert ws["I7"].value == "=I8/2"


def test_report_undefined_percent_is_not_zero():
    students = [student(1, "Антонов Антон")]
    events = [event(1, "2026-09-03", kind="free")]
    stats = compute_group_stats(students, events, {1: {1}}, [])
    assert stats[0].percent is None
    wb = _book(
        build_report_xlsx(
            org_title="Посещаемость групп",
            period_title="Сентябрь 2026",
            period_from=date(2026, 9, 3),
            period_to=date(2026, 9, 3),
            stats=stats,
            events=events,
        )
    )
    ws = wb["Сентябрь 2026"]
    assert ws["C5"].value == "—"
    assert ws["I4"].value is None
    assert ws["I5"].value == 1
    assert ws["J4"].value == "=SUM(I4:I4)"


def test_report_percent_formulas_and_interim_mark():
    students = [student(1, "Антонов Антон"), student(2, "Борисов Борис")]
    events = [event(1, "2026-09-01"), event(2, "2026-09-03", kind="free"), event(3, "2026-09-05")]
    stats = compute_group_stats(students, events, {1: {1, 2}, 2: {1}, 3: {1}}, [])
    by_name = {item.student.full_name: item.percent for item in stats}
    assert by_name["Антонов Антон"] == 150.0
    assert by_name["Борисов Борис"] == 50.0
    wb = _book(
        build_report_xlsx(
            org_title="Посещаемость групп",
            period_title="Сентябрь 2026",
            period_from=date(2026, 9, 1),
            period_to=date(2026, 9, 5),
            stats=stats,
            events=events,
            interim=True,
        )
    )
    ws = wb["Сентябрь 2026"]
    assert ws["A1"].value == "промежуточный"
    assert (ws["I5"].value, ws["J5"].value, ws["K5"].value) == (1, 1, 1)
    assert (ws["I6"].value, ws["J6"].value, ws["K6"].value) == (1, None, None)
    assert len(ws.conditional_formatting._cf_rules) == 1


def test_same_day_lessons_share_a_header():
    students = [student(1, "Антонов Антон")]
    events = [
        event(1, "2026-09-01", title="Шахарит", time="08:00"),
        event(2, "2026-09-01", title="Иудаика", time="10:00"),
    ]
    stats = compute_group_stats(students, events, {1: {1}, 2: {1}}, [])
    wb = _book(
        build_report_xlsx(
            org_title="Посещаемость групп",
            period_title="Сентябрь 2026",
            period_from=date(2026, 9, 1),
            period_to=date(2026, 9, 1),
            stats=stats,
            events=events,
        )
    )
    ws = wb["Сентябрь 2026"]
    assert {str(item) for item in ws.merged_cells.ranges} >= {"I1:J1", "I2:J2"}
    assert ws["I1"].value == "вторник"
    assert ws["I2"].value == 1
    assert (ws["I3"].value, ws["J3"].value) == ("Шахарит", "Иудаика")
    assert (ws["I5"].value, ws["J5"].value) == (1, 1)


def test_late_joiner_uses_own_denominator():
    students = [student(1, "Антонов Антон"), student(2, "Борисов Борис", joined="2026-09-05")]
    events = [event(1, "2026-09-01"), event(2, "2026-09-05")]
    stats = compute_group_stats(students, events, {2: {2}}, [])
    wb = _book(
        build_report_xlsx(
            org_title="Посещаемость групп",
            period_title="Сентябрь 2026",
            period_from=date(2026, 9, 1),
            period_to=date(2026, 9, 5),
            stats=stats,
            events=events,
        )
    )
    ws = wb["Сентябрь 2026"]
    assert ws["C5"].value == "=D5/$K$4"
    assert (ws["I5"].value, ws["J5"].value) == (None, None)
    assert ws["C6"].value == "=D6/1"
    assert (ws["I6"].value, ws["J6"].value) == (None, 1)


def test_archive_has_one_sheet_per_period():
    students = [student(1, "Антонов Антон")]
    closed_events = [event(1, "2026-09-01", closes=True)]
    open_events = [event(2, "2026-10-02", title="Семинар")]
    closed = compute_group_stats(students, closed_events, {1: {1}}, [])
    opened = compute_group_stats(students, open_events, {}, [])
    wb = _book(
        build_periods_xlsx(
            org_title="Посещаемость групп",
            sections=[
                PeriodSection("Сентябрь: 2026", date(2026, 9, 1), date(2026, 9, 1), closed, closed_events),
                PeriodSection(
                    "Текущий: Октябрь 2026",
                    date(2026, 10, 2),
                    date(2026, 10, 2),
                    opened,
                    open_events,
                    interim=True,
                ),
            ],
        )
    )
    assert len(wb.sheetnames) == 2
    assert all(re.search(r"[\\*?:/\[\]]", name) is None for name in wb.sheetnames)
    assert wb.sheetnames == ["Сентябрь 2026", "Текущий Октябрь 2026"]
    assert "промежуточный" in wb[wb.sheetnames[1]]["A1"].value
    assert wb[wb.sheetnames[0]].cell(row=5, column=9).value == 1
    assert wb[wb.sheetnames[1]].cell(row=3, column=9).value == "Семинар"
